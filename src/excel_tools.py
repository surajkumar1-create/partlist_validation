"""
Tools that the AI agent calls to get EXACT answers from the Excel file.
The model never counts or reads rows itself: it calls these functions and
explains the result. Every function returns small, JSON-friendly data.
"""
import difflib
import hashlib
import io
import re
import threading
from collections import Counter

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# sheet name -> (header row, first data row). Other sheets: header is guessed, data starts below it.
SHEET_DEFAULTS = {"ZaragozaTH,PF,BS": (4, 7)}
MAX_COLS = 300
BLANK_STOP = 50
EXCEL_ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}

_LOCK = threading.RLock()
_STATE = {"key": None, "bytes": None, "wb": None, "tables": {}}


def set_file(file_bytes: bytes) -> None:
    """Called by the app once per uploaded file (not an AI tool)."""
    key = hashlib.sha1(file_bytes).hexdigest()
    with _LOCK:
        if _STATE["key"] != key:
            _STATE.update(key=key, bytes=file_bytes, wb=None, tables={})


def _wb():
    with _LOCK:
        if _STATE["wb"] is None:
            _STATE["wb"] = load_workbook(io.BytesIO(_STATE["bytes"]), read_only=True,
                                         data_only=True, keep_links=False)
        return _STATE["wb"]


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and v != v) or (isinstance(v, str) and v.replace("\xa0", " ").strip() == "")


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    return str(v)


def _resolve_sheet(sheet: str) -> str:
    names = _wb().sheetnames
    if sheet in names:
        return sheet
    low = {n.lower().replace(" ", ""): n for n in names}
    key = sheet.lower().replace(" ", "")
    if key in low:
        return low[key]
    close = [n for n in names if key in n.lower().replace(" ", "")]
    if len(close) == 1:
        return close[0]
    raise ValueError(f"Sheet '{sheet}' not found. Available sheets: {names}")


def _table(sheet: str):
    """DataFrame (index = Excel row, columns = letters) + headers dict, cached per sheet."""
    sheet = _resolve_sheet(sheet)
    with _LOCK:
        if sheet in _STATE["tables"]:
            return _STATE["tables"][sheet]
        ws = _wb()[sheet]
        ncols = min(ws.max_column or MAX_COLS, MAX_COLS)
        if sheet in SHEET_DEFAULTS:
            hdr_row, first = SHEET_DEFAULTS[sheet]
        else:
            best, best_n = 1, 0
            for i, r in enumerate(ws.iter_rows(min_row=1, max_row=30, max_col=ncols, values_only=True), 1):
                n = sum(1 for c in r if isinstance(c, str) and c.strip())
                if n > best_n:
                    best, best_n = i, n
            hdr_row, first = best, best + 1
        header = next(ws.iter_rows(min_row=hdr_row, max_row=hdr_row, max_col=ncols, values_only=True), ())
        header = list(header) + [None] * (ncols - len(header))
        data, rows, blank_run = [], [], 0
        for rn, row in enumerate(ws.iter_rows(min_row=first, max_col=ncols, values_only=True), start=first):
            if all(_blank(c) for c in row):
                blank_run += 1
                if blank_run >= BLANK_STOP:
                    break
                continue
            blank_run = 0
            data.append(list(row) + [None] * (ncols - len(row)))
            rows.append(rn)
        letters = [get_column_letter(i + 1) for i in range(ncols)]
        df = pd.DataFrame(data, columns=letters, index=rows, dtype=object)
        keep = 0
        for i, l in enumerate(letters):
            if not _blank(header[i]) or df[l].map(lambda v: not _blank(v)).any():
                keep = i + 1
        df = df[letters[:keep]]
        headers = {l: (re.sub(r"\s+", " ", str(header[i])).strip() if not _blank(header[i]) else "")
                   for i, l in enumerate(letters[:keep])}
        _STATE["tables"][sheet] = (sheet, df, headers, hdr_row, first)
        return _STATE["tables"][sheet]


def _column(headers: dict, spec: str) -> str:
    s = str(spec).strip()
    for l, h in headers.items():                       # exact header text
        if h and h.lower() == s.lower():
            return l
    if s.upper() in headers:                           # Excel letter
        return s.upper()
    close = [l for l, h in headers.items() if h and s.lower() in h.lower()]
    if len(close) == 1:
        return close[0]
    names = {l: h.lower() for l, h in headers.items() if h}      # tolerate typos, e.g. "Commonality"
    fuzzy = difflib.get_close_matches(s.lower(), list(names.values()), n=2, cutoff=0.85)
    if len(fuzzy) == 1 and list(names.values()).count(fuzzy[0]) == 1:
        return next(l for l, h in names.items() if h == fuzzy[0])
    options = {l: h for l, h in headers.items() if h}
    raise ValueError(f"Column '{spec}' not found or ambiguous ({close}). Columns: {options}")


def _ignore(df, col, ignore_values: str, sheet: str = ""):
    """Optional: drop rows whose value in this column equals one of ignore_values.
    By default nothing is dropped, so answers match the sheet exactly."""
    vals = {v.strip().lower() for v in ignore_values.split(";") if v.strip()}
    if not vals:
        return df
    return df[~df[col].map(lambda v: _s(v).strip().lower() in vals)]


# ------------------------------------------------------------------ AI tools
def workbook_overview() -> dict:
    """List every sheet in the workbook with its visibility and, for the main data sheets, the
    detected header row, first data row, number of data rows and column count.
    Use this when the user asks what is in the workbook."""
    wb = _wb()
    out = []
    for ws in wb.worksheets:
        info = {"sheet": ws.title, "visibility": ws.sheet_state}
        if ws.title in SHEET_DEFAULTS:
            sheet, df, headers, hr, fr = _table(ws.title)
            info.update(header_row=hr, first_data_row=fr, data_rows=len(df), columns=len(headers))
        out.append(info)
    return {"sheets": out}


def column_summary(sheet: str, column: str, ignore_values: str = "") -> dict:
    """Profile ONE column: how many cells are filled/empty, how many distinct values, Excel errors
    (#N/A etc.), cells with extra spaces, and the 15 most frequent values with counts.

    Args:
        sheet: Sheet name, e.g. "ZaragozaTH,PF,BS".
        column: Column header text or Excel letter, e.g. "Part number" or "F".
        ignore_values: Values to exclude, separated by ';'. Leave empty (default) to count every row exactly as in the sheet. Only fill this if the user explicitly asks to exclude something.
    """
    sheet, df, headers, hr, fr = _table(sheet)
    c = _column(headers, column)
    d = _ignore(df, c, ignore_values, sheet)
    vals = [_s(v).strip() for v in d[c] if not _blank(v)]
    cnt = Counter(vals)
    return {"sheet": sheet, "column": c, "header": headers[c], "rows_considered": len(d),
            "ignored_rows": len(df) - len(d), "filled": len(vals), "empty": len(d) - len(vals),
            "distinct_values": len(cnt), "excel_errors": sum(1 for v in vals if v in EXCEL_ERRORS),
            "cells_with_extra_spaces": sum(1 for v in d[c] if isinstance(v, str)
                                           and v != " ".join(v.replace("\xa0", " ").split())),
            "top_values": [{"value": k, "count": n} for k, n in cnt.most_common(15)]}


def count_duplicates(sheet: str, column: str, ignore_values: str = "") -> dict:
    """Count duplicate values in ONE column and list the 15 most repeated, with the exact cell
    addresses of each occurrence (up to 8). Empty cells are ignored.

    Args:
        sheet: Sheet name, e.g. "ZaragozaTH,PF,BS".
        column: Column header text or Excel letter, e.g. "Part number" or "F".
        ignore_values: Values to exclude, separated by ';'. Leave empty (default) to count every row exactly as in the sheet. Only fill this if the user explicitly asks to exclude something.
    """
    sheet, df, headers, hr, fr = _table(sheet)
    c = _column(headers, column)
    kept = _ignore(df, c, ignore_values, sheet)
    d = kept[~kept[c].map(_blank)]
    keys = d[c].map(lambda v: _s(v).strip())
    counts = keys.value_counts()
    dup = counts[counts > 1]
    examples = []
    for val, n in dup.head(15).items():
        rows = list(keys[keys == val].index[:8])
        examples.append({"value": val, "occurrences": int(n), "cells": [f"{c}{r}" for r in rows]})
    return {"sheet": sheet, "column": c, "header": headers[c], "rows_considered": len(d),
            "ignored_rows": len(df) - len(kept),
            "distinct_values": int(counts.size), "values_that_repeat": int(dup.size),
            "rows_involved_in_duplicates": int(dup.sum()), "extra_duplicate_rows": int((dup - 1).sum()),
            "most_repeated": examples}


def find_rows(sheet: str, column: str, value: str, match: str = "equals", max_rows: int = 15) -> dict:
    """Find rows where a column equals (or contains) a value. Returns the exact Excel row numbers,
    the cell address of the match and the first non-empty cells of each row.

    Args:
        sheet: Sheet name, e.g. "ZaragozaTH,PF,BS".
        column: Column header text or Excel letter, e.g. "Part number" or "F".
        value: Text to look for (case-insensitive, spaces trimmed).
        match: "equals" or "contains".
        max_rows: Maximum rows to return (the total match count is always given).
    """
    sheet, df, headers, hr, fr = _table(sheet)
    c = _column(headers, column)
    v = value.strip().lower()
    keys = df[c].map(lambda x: _s(x).strip().lower())
    hit = keys.str.contains(v, regex=False) if match == "contains" else keys == v
    found = df[hit]
    rows = []
    for rn, row in found.head(min(int(max_rows), 30)).iterrows():
        cells = {f"{l}{rn}": _s(row[l])[:60] for l in headers if not _blank(row[l])}
        rows.append({"row": int(rn), "match_cell": f"{c}{rn}", "cells": dict(list(cells.items())[:14])})
    return {"sheet": sheet, "column": c, "header": headers[c], "total_matches": int(hit.sum()),
            "returned": len(rows), "rows": rows}