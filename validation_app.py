
"""
Partlist Validation App
-----------------------
Streamlit UI for deterministic Excel rule validation.

Features:
- Upload a partlist workbook and an optional rules workbook.
- Validate the workbook using src/rule_checker.py.
- Display an overall dashboard.
- Select a column and see only its results.
- Edit failed values and revalidate.
- Open a selected failed cell in desktop Excel.
- Download the corrected workbook and validation report.
"""

from __future__ import annotations

import io
import math
import os
import re
import tempfile
import uuid
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src import excel_tools as T
from src import rule_checker as R
from src.ui_helpers import selected_row_index


# =========================================================
# 1. CONFIGURATION
# =========================================================

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_RULES_PATH = BASE_DIR / "Stated_Rule_List.xlsx"
MAIN_SHEET = "ZaragozaTH,PF,BS"

st.set_page_config(
    page_title="Partlist Validation",
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    .block-container {
        max-width: 1400px;
        padding-top: 1.2rem;
        padding-bottom: 2rem;
    }
    h1 {
        font-size: 1.9rem !important;
    }
    [data-testid="stMetric"] {
        background: var(--secondary-background-color);
        border: 1px solid rgba(128, 140, 160, 0.25);
        border-radius: 10px;
        padding: 12px;
    }
    div[data-testid="stDataFrame"],
    div[data-testid="stDataEditor"] {
        border-radius: 8px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("Partlist Validation")
st.caption(
    "Validate your workbook, investigate issues by column, "
    "correct values, and download the results."
)


# =========================================================
# 2. GENERAL HELPERS
# =========================================================

def is_blank(value: Any) -> bool:
    """Return True when a value is empty or NaN."""
    if value is None:
        return True

    if isinstance(value, float) and math.isnan(value):
        return True

    return isinstance(value, str) and value.strip() == ""


def values_are_equal(a: Any, b: Any) -> bool:
    """Compare values without treating blank cells as different."""
    if is_blank(a) and is_blank(b):
        return True

    if (
        isinstance(a, (int, float))
        and not isinstance(a, bool)
        and isinstance(b, (int, float))
        and not isinstance(b, bool)
    ):
        return float(a) == float(b)

    return str(a).strip() == str(b).strip()


def safe_text(value: Any) -> str:
    """Convert a value to readable text."""
    if is_blank(value):
        return ""

    return str(value)


def normalize_column(value: Any) -> str:
    """Normalize a column identifier for comparison."""
    return str(value).strip().upper()


def status_counts(rule_summary: pd.DataFrame) -> dict[str, int]:
    """Count PASS, FAIL and REVIEW rule results."""
    counts = {
        "PASS": 0,
        "FAIL": 0,
        "REVIEW": 0,
    }

    if rule_summary.empty or "status" not in rule_summary.columns:
        return counts

    statuses = (
        rule_summary["status"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
    )

    for status in counts:
        counts[status] = int((statuses == status).sum())

    return counts


def suggested_correction(
    rule_ids: str,
    message: str,
) -> str:
    """Provide a general correction hint based on a finding."""
    ids = set(re.findall(r"R\d{3}", rule_ids or ""))
    text = (message or "").casefold()

    if "duplicate" in text or "R001" in ids:
        return (
            "Check the duplicate rows and enter the approved "
            "part number from the source data."
        )

    if ids.intersection({"R029", "R030", "R031"}):
        return "Enter the correct positive whole number."

    if "whole number" in text or "not a number" in text:
        return "Enter a valid whole number according to the rule."

    if ids.intersection({"R025", "R026", "R027", "R028"}):
        return (
            "Check whether the corresponding unit applies, "
            "then enter the required value."
        )

    if "blank" in text or "required" in text:
        return "Enter the required value from the source data."

    if ids.intersection({"R015", "R016"}):
        return (
            "Verify the SILS value and its exemption conditions "
            "against the source data."
        )

    return (
        "Check the rule description and source data before "
        "changing this value."
    )


# =========================================================
# 3. BUILD FAILED-CELL TABLE
# =========================================================

def build_failed_cells(result: dict[str, Any]) -> pd.DataFrame:
    """
    Group findings by cell.

    Expected finding fields include:
    cell, sheet, excel_row, column, header,
    actual_value, rule_id, message.
    """

    grouped: dict[str, dict[str, Any]] = {}

    for finding in result.get("findings", []):
        cell = safe_text(finding.get("cell"))

        if not cell:
            continue

        if cell not in grouped:
            grouped[cell] = {
                "Sheet": finding.get(
                    "sheet",
                    result.get("sheet", ""),
                ),
                "Excel Row": finding.get("excel_row", ""),
                "Cell": cell,
                "Column": finding.get("column", ""),
                "Header": finding.get("header", ""),
                "Current Value": finding.get("actual_value", ""),
                "Rule IDs": [],
                "Issue": [],
            }

        rule_id = safe_text(finding.get("rule_id"))
        message = safe_text(finding.get("message"))

        if rule_id and rule_id not in grouped[cell]["Rule IDs"]:
            grouped[cell]["Rule IDs"].append(rule_id)

        if message and message not in grouped[cell]["Issue"]:
            grouped[cell]["Issue"].append(message)

    rows = []

    for item in grouped.values():
        item["Rule IDs"] = ", ".join(item["Rule IDs"])
        item["Issue"] = " | ".join(item["Issue"])

        item["Required Correction"] = suggested_correction(
            item["Rule IDs"],
            item["Issue"],
        )

        item["Corrected Value"] = item["Current Value"]
        rows.append(item)

    columns = [
        "Sheet",
        "Excel Row",
        "Cell",
        "Column",
        "Header",
        "Current Value",
        "Rule IDs",
        "Issue",
        "Required Correction",
        "Corrected Value",
    ]

    return pd.DataFrame(rows, columns=columns)


# =========================================================
# 4. EDIT EXCEL CELLS IN MEMORY
# =========================================================

def convert_excel_value(
    new_value: Any,
    old_value: Any,
) -> Any:
    """Convert an edited value to a suitable Excel value."""

    if is_blank(new_value):
        return None

    if isinstance(old_value, bool):
        if isinstance(new_value, str):
            value = new_value.strip().casefold()

            if value in {"true", "yes", "1"}:
                return True

            if value in {"false", "no", "0"}:
                return False

        return bool(new_value)

    if isinstance(old_value, int) and not isinstance(old_value, bool):
        try:
            number = float(str(new_value).strip())

            if number.is_integer():
                return int(number)

            return number
        except (TypeError, ValueError):
            return new_value

    if isinstance(old_value, float):
        try:
            return float(str(new_value).strip())
        except (TypeError, ValueError):
            return new_value

    if isinstance(new_value, (int, float)) and not isinstance(
        new_value, bool
    ):
        return new_value

    return str(new_value)


def write_cells(
    workbook_bytes: bytes,
    filename: str,
    sheet_name: str,
    edits: list[tuple[str, Any]],
) -> bytes:
    """Apply edits and return the updated workbook as bytes."""

    keep_vba = filename.lower().endswith(".xlsm")

    workbook = load_workbook(
        io.BytesIO(workbook_bytes),
        keep_vba=keep_vba,
        data_only=False,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet {sheet_name!r} was not found."
            )

        worksheet = workbook[sheet_name]

        for address, new_value in edits:
            old_value = worksheet[address].value
            worksheet[address] = convert_excel_value(
                new_value,
                old_value,
            )

        output = io.BytesIO()
        workbook.save(output)

        return output.getvalue()

    finally:
        workbook.close()


# =========================================================
# 5. OPEN SELECTED CELL IN DESKTOP EXCEL
# =========================================================

def open_corrected_workbook_at_cell(
    workbook_bytes: bytes,
    filename: str,
    sheet_name: str,
    cell_address: str,
) -> Path:
    """
    Save a temporary workbook and open the requested cell in Excel.

    Requires Windows, Microsoft Excel and pywin32.
    """

    if os.name != "nt":
        raise RuntimeError(
            "Opening a selected cell in desktop Excel "
            "requires Windows."
        )

    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError(
            "Install pywin32 using: python -m pip install pywin32"
        ) from exc

    suffix = Path(filename).suffix.lower() or ".xlsx"
    stem = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        Path(filename).stem,
    ) or "partlist"

    target = (
        Path(tempfile.gettempdir())
        / f"{stem}_corrected_{uuid.uuid4().hex[:8]}{suffix}"
    )

    target.write_bytes(workbook_bytes)

    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = True

    workbook = excel.Workbooks.Open(str(target.resolve()))

    try:
        worksheet = workbook.Worksheets(sheet_name)
        worksheet.Activate()
        excel.Goto(worksheet.Range(cell_address), True)
        worksheet.Range(cell_address).Select

        try:
            excel.WindowState = -4137
        except Exception:
            pass

    except Exception:
        workbook.Close(SaveChanges=False)
        raise

    return target


# =========================================================
# 6. BUILD VALIDATION REPORT
# =========================================================

def build_report(
    result: dict[str, Any],
    failed: pd.DataFrame,
) -> bytes:
    """Create an Excel report with Summary, Rule Results and Failed Cells."""

    output = io.BytesIO()
    workbook = Workbook()

    # Summary worksheet
    summary_ws = workbook.active
    summary_ws.title = "Summary"

    summary_ws.append(["Metric", "Value"])

    summary_metrics = [
        ("Worksheet", result.get("sheet", "")),
        (
            "Rows checked",
            result.get("data_rows_checked", 0),
        ),
        (
            "Rules loaded",
            result.get("rules_loaded", 0),
        ),
        (
            "Rules passed",
            result.get("rules_passed", 0),
        ),
        (
            "Rules failed",
            result.get("rules_failed", 0),
        ),
        (
            "Rules needing review",
            result.get("rules_needing_review", 0),
        ),
        ("Failed cells", len(failed)),
    ]

    for metric in summary_metrics:
        summary_ws.append(list(metric))

    # Rule Results worksheet
    rule_ws = workbook.create_sheet("Rule Results")

    rule_ws.append(
        [
            "Rule ID",
            "Column",
            "Header",
            "Rule",
            "Status",
            "Message",
            "Sheet",
        ]
    )

    for item in result.get("rule_summary", []):
        rule_ws.append(
            [
                item.get("rule_id", ""),
                item.get("column", ""),
                item.get("header", ""),
                item.get("rule", ""),
                item.get("status", ""),
                item.get("message", ""),
                item.get("sheet", ""),
            ]
        )

    # Failed Cells worksheet
    failed_ws = workbook.create_sheet("Failed Cells")
    failed_ws.append(list(failed.columns))

    for row in failed.itertuples(index=False, name=None):
        failed_ws.append(
            [
                None if is_blank(value) else value
                for value in row
            ]
        )

    # Review worksheet
    review_ws = workbook.create_sheet("Review")

    review_ws.append(
        [
            "Rule ID",
            "Column",
            "Header",
            "Rule",
            "Status",
            "Reason",
            "Sheet",
        ]
    )

    for item in result.get("rule_summary", []):
        if safe_text(item.get("status")).upper() == "REVIEW":
            review_ws.append(
                [
                    item.get("rule_id", ""),
                    item.get("column", ""),
                    item.get("header", ""),
                    item.get("rule", ""),
                    item.get("status", ""),
                    item.get("message", ""),
                    item.get("sheet", ""),
                ]
            )

    # Apply report formatting
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

        if worksheet.max_row > 1:
            worksheet.auto_filter.ref = worksheet.dimensions

        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter

            max_length = max(
                (
                    len(str(cell.value))
                    if cell.value is not None
                    else 0
                    for cell in column_cells[:1000]
                ),
                default=12,
            )

            worksheet.column_dimensions[letter].width = min(
                max(max_length + 2, 12),
                55,
            )

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    wrap_text=True,
                    vertical="top",
                )

    workbook.save(output)
    workbook.close()

    return output.getvalue()


# =========================================================
# 7. SIDEBAR: FILE UPLOADS
# =========================================================

with st.sidebar:
    st.header("Input files")

    uploaded_file = st.file_uploader(
        "Upload partlist workbook",
        type=["xlsx", "xlsm"],
        help="Choose the workbook you want to validate.",
    )

    uploaded_rules = st.file_uploader(
        "Upload rules workbook (optional)",
        type=["xlsx"],
        help=(
            "If omitted, the app uses Stated_Rule_List.xlsx "
            "from the project folder."
        ),
    )

    st.divider()

    st.caption(
        "Validation is performed by Python rules. "
        "No AI model is called."
    )


# =========================================================
# 8. WAIT FOR A WORKBOOK
# =========================================================

if uploaded_file is None:
    st.info(
        "Upload a partlist workbook using the sidebar "
        "to start validation."
    )

    st.markdown(
        """
        **Workflow**

        1. Upload your partlist.
        2. Review the overall dashboard.
        3. Select a column to inspect its results.
        4. Correct failed values and revalidate.
        5. Download the corrected workbook and report.
        """
    )

    st.stop()


# =========================================================
# 9. LOAD WORKBOOK AND RULES
# =========================================================

original_filename = uploaded_file.name
original_bytes = uploaded_file.getvalue()

if uploaded_rules is not None:
    rules_bytes = uploaded_rules.getvalue()
else:
    if not DEFAULT_RULES_PATH.exists():
        st.error(
            "Stated_Rule_List.xlsx was not found in the project "
            "folder. Upload the rules workbook or place it beside "
            "validation_app.py."
        )
        st.stop()

    rules_bytes = DEFAULT_RULES_PATH.read_bytes()


# Reset edits when the uploaded workbook changes.
import hashlib

file_hash = hashlib.sha256(original_bytes).hexdigest()

if st.session_state.get("source_file_hash") != file_hash:
    st.session_state["source_file_hash"] = file_hash
    st.session_state["working_bytes"] = original_bytes
    st.session_state["working_filename"] = original_filename
    st.session_state.pop("last_opened_excel_selection", None)


working_bytes = st.session_state["working_bytes"]
working_filename = st.session_state["working_filename"]

working_hash = hashlib.sha256(working_bytes).hexdigest()


# Read available worksheet names.
try:
    source_workbook = load_workbook(
        io.BytesIO(working_bytes),
        read_only=True,
        data_only=True,
        keep_vba=working_filename.lower().endswith(".xlsm"),
    )

    try:
        sheet_names = list(source_workbook.sheetnames)
    finally:
        source_workbook.close()

except Exception as exc:
    st.error(f"Could not read the workbook: {exc}")
    st.stop()


if not sheet_names:
    st.error("The workbook does not contain any worksheets.")
    st.stop()


preferred_sheet = (
    MAIN_SHEET
    if MAIN_SHEET in sheet_names
    else sheet_names[0]
)

selected_sheet = st.selectbox(
    "Worksheet to validate",
    options=sheet_names,
    index=sheet_names.index(preferred_sheet),
    key=f"worksheet_{file_hash}",
)


# =========================================================
# 10. RUN VALIDATION
# =========================================================

try:
    # These calls use the existing project module interfaces.
    T.set_file(working_bytes)
    R.set_rules_file(rules_bytes)

    result = R.check_all_rules(selected_sheet)

except Exception as exc:
    st.error(
        "Validation could not run. Check the interfaces in "
        "src/excel_tools.py and src/rule_checker.py."
    )

    st.exception(exc)
    st.stop()


if not isinstance(result, dict):
    st.error(
        "check_all_rules() must return a dictionary containing "
        "findings and rule_summary."
    )
    st.stop()


failed = build_failed_cells(result)

rule_summary = pd.DataFrame(
    result.get("rule_summary", [])
)

counts = status_counts(rule_summary)


# =========================================================
# 11. OVERALL DASHBOARD
# =========================================================

st.divider()
st.subheader("Overall dashboard")

metric1, metric2, metric3, metric4 = st.columns(4)

metric1.metric(
    "Rows checked",
    result.get("data_rows_checked", 0),
)

metric2.metric(
    "Failed cells",
    len(failed),
)

metric3.metric(
    "Rules failed",
    result.get("rules_failed", counts["FAIL"]),
)

metric4.metric(
    "Needs review",
    result.get("rules_needing_review", counts["REVIEW"]),
)

st.caption(
    f"Worksheet: {selected_sheet}  ·  "
    f"Rules loaded: {result.get('rules_loaded', len(rule_summary))}  ·  "
    f"Rules passed: {result.get('rules_passed', counts['PASS'])}"
)

if result.get("rules_needing_review", counts["REVIEW"]) > 0:
    st.warning(
        "Some rules need manual review. Review these findings "
        "before treating the workbook as fully validated."
    )

with st.expander("View overall rule summary"):
    if rule_summary.empty:
        st.info("No rule summary was returned by the rule checker.")
    else:
        st.dataframe(
            rule_summary,
            use_container_width=True,
            hide_index=True,
        )


# =========================================================
# 12. COLUMN SELECTOR AND COLUMN-SPECIFIC RESULTS
# =========================================================

st.divider()
st.subheader("Column-wise validation results")

if failed.empty:
    st.success("No failed cells were found in the workbook.")

    st.info(
        "The column selector appears when the rule checker "
        "returns failed cells."
    )

else:
    # Create a dropdown from the failed cells.
    column_options = (
        failed[["Column", "Header"]]
        .fillna("")
        .drop_duplicates()
        .copy()
    )

    column_options["Label"] = (
        column_options["Column"].astype(str)
        + " — "
        + column_options["Header"].astype(str)
    )

    selected_label = st.selectbox(
        "Select a column",
        options=column_options["Label"].tolist(),
        key=f"column_selector_{working_hash}_{selected_sheet}",
    )

    selected_option = column_options.loc[
        column_options["Label"] == selected_label
    ].iloc[0]

    selected_column = selected_option["Column"]

    # Filter the failed-cell table.
    column_failed = failed[
        failed["Column"].apply(normalize_column)
        == normalize_column(selected_column)
    ].copy()

    # Filter rule results for this column.
    if (
        not rule_summary.empty
        and "column" in rule_summary.columns
    ):
        column_rule_summary = rule_summary[
            rule_summary["column"].apply(normalize_column)
            == normalize_column(selected_column)
        ].copy()
    else:
        column_rule_summary = pd.DataFrame()

    st.markdown(f"### Results: {selected_label}")

    # Column-specific metrics.
    c1, c2, c3 = st.columns(3)

    c1.metric(
        "Failed cells",
        len(column_failed),
    )

    column_counts = status_counts(column_rule_summary)

    c2.metric(
        "Failed rules",
        column_counts["FAIL"],
    )

    c3.metric(
        "Rules needing review",
        column_counts["REVIEW"],
    )

    # Rule details for this column only.
    with st.expander(
        "Rule results for selected column",
        expanded=True,
    ):
        if column_rule_summary.empty:
            st.info(
                "No rule summary entries matched this column. "
                "The failed-cell findings are shown below."
            )
        else:
            display_columns = [
                column
                for column in [
                    "rule_id",
                    "column",
                    "header",
                    "status",
                    "message",
                ]
                if column in column_rule_summary.columns
            ]

            st.dataframe(
                column_rule_summary[display_columns].rename(
                    columns={
                        "rule_id": "Rule ID",
                        "column": "Column",
                        "header": "Field",
                        "status": "Status",
                        "message": "Result",
                    }
                ),
                use_container_width=True,
                hide_index=True,
            )

    # Show only selected column's failed cells.
    st.markdown("#### Failed cells in this column")

    if column_failed.empty:
        st.success("No failed cells in the selected column.")

    else:
        st.caption(
            "Select a cell to open the workbook at that location "
            "in desktop Excel. You can also edit its corrected value."
        )

        navigation_columns = [
            column
            for column in [
                "Cell",
                "Header",
                "Current Value",
                "Rule IDs",
                "Issue",
            ]
            if column in column_failed.columns
        ]

        navigation_event = st.dataframe(
            column_failed[navigation_columns],
            key=(
                f"cell_navigation_{working_hash}_"
                f"{selected_sheet}_{selected_column}"
            ),
            use_container_width=True,
            hide_index=True,
            on_select="rerun",
            selection_mode="single-cell",
        )

        # Open the selected failed cell in Excel.
        selected_row = selected_row_index(
            getattr(
                navigation_event,
                "selection",
                None,
            )
        )

        if (
            selected_row is not None
            and 0 <= selected_row < len(column_failed)
        ):
            selected_cell = str(
                column_failed.iloc[selected_row]["Cell"]
            )

            selection_token = (
                f"{working_hash}|"
                f"{selected_sheet}|"
                f"{selected_cell}"
            )

            if (
                st.session_state.get(
                    "last_opened_excel_selection"
                )
                != selection_token
            ):
                try:
                    open_corrected_workbook_at_cell(
                        working_bytes,
                        working_filename,
                        selected_sheet,
                        selected_cell,
                    )

                    st.session_state[
                        "last_opened_excel_selection"
                    ] = selection_token

                    st.toast(
                        f"Opened {selected_sheet}!{selected_cell}"
                    )

                except Exception as exc:
                    st.warning(
                        f"Excel navigation was unavailable: {exc}"
                    )

        # Edit corrected values for the selected column.
        st.markdown("#### Correct values")

        edited_column = st.data_editor(
            column_failed,
            key=(
                f"failed_editor_{working_hash}_"
                f"{selected_sheet}_{selected_column}"
            ),
            use_container_width=True,
            hide_index=True,
            disabled=[
                column
                for column in column_failed.columns
                if column != "Corrected Value"
            ],
            column_config={
                "Corrected Value": st.column_config.TextColumn(
                    "Corrected Value",
                    help="Enter the correct value for this Excel cell.",
                ),
                "Required Correction": st.column_config.TextColumn(
                    "Suggested correction",
                    width="large",
                ),
                "Issue": st.column_config.TextColumn(
                    "Validation issue",
                    width="large",
                ),
            },
        )

        if st.button(
            "Apply edits and revalidate",
            type="primary",
            key=(
                f"apply_edits_{working_hash}_"
                f"{selected_sheet}_{selected_column}"
            ),
        ):
            edits = []

            for _, row in edited_column.iterrows():
                if not values_are_equal(
                    row["Corrected Value"],
                    row["Current Value"],
                ):
                    edits.append(
                        (
                            str(row["Cell"]),
                            row["Corrected Value"],
                        )
                    )

            if not edits:
                st.info("No changes detected.")

            else:
                try:
                    updated_bytes = write_cells(
                        working_bytes,
                        working_filename,
                        selected_sheet,
                        edits,
                    )

                    st.session_state["working_bytes"] = updated_bytes
                    st.session_state.pop(
                        "last_opened_excel_selection",
                        None,
                    )

                    st.success(
                        f"Applied {len(edits)} edit(s). "
                        "The workbook will now be revalidated."
                    )

                    st.rerun()

                except Exception as exc:
                    st.error(f"Could not apply edits: {exc}")


# =========================================================
# 13. EXPORT CORRECTED WORKBOOK AND VALIDATION REPORT
# =========================================================

st.divider()
st.subheader("Download results")

try:
    report_bytes = build_report(result, failed)

    output_stem = Path(working_filename).stem
    suffix = Path(working_filename).suffix.lower() or ".xlsx"

    if suffix == ".xlsm":
        workbook_mime = (
            "application/vnd.ms-excel.sheet.macroEnabled.12"
        )
    else:
        workbook_mime = (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

    download1, download2 = st.columns(2)

    with download1:
        st.download_button(
            "Download corrected partlist",
            data=working_bytes,
            file_name=f"{output_stem}_corrected{suffix}",
            mime=workbook_mime,
            use_container_width=True,
        )

    with download2:
        st.download_button(
            "Download validation report",
            data=report_bytes,
            file_name=f"{output_stem}_validation_report.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True,
        )

except Exception as exc:
    st.error(f"Could not create the downloadable report: {exc}")


# =========================================================
# 14. NOTES
# =========================================================

with st.expander("Validation notes"):
    notes = result.get("notes", [])

    if notes:
        for note in notes:
            st.write(f"- {note}")
    else:
        st.write(
            "The app uses the existing Python rule checker. "
            "Review conditional rules and source data before "
            "approving corrections."
        )
