from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from src import excel_tools as T
from src import rule_checker as R

ROOT = Path(__file__).resolve().parents[1]
RULES_PATH = ROOT / "Stated_Rule_List.xlsx"
SHEET = "ZaragozaTH,PF,BS"

# Enough columns to exercise the uploaded 31-rule workbook.
HEADERS = {
    "F": "Part number", "G": "French Designation", "H": "Designation English",
    "I": "Family", "J": "Code Function", "M": "Fastners/Part/DOT/ Consumables",
    "N": "Part Route", "O": "P1H", "P": "O1H", "Q": "Commanality",
    "Y": "COFOR seller", "Z": "Seller's company name", "AA": "COFOR shipper",
    "AB": "Country", "AF": "Packaging Remarks", "AU": "UC code",
    "AV": "UC Capacity", "AY": "Type", "AZ": "UC length", "BA": "UC width",
    "BB": "UC height", "BC": "Std/Non Std Packaging", "BD": "Reusable / Disposable",
    "BE": "Substance", "BO": "P1H required", "BT": "O1H required",
    "BV": "% Utile P1H", "BW": "% Utile O1H", "BX": "RPV - P1H",
    "BY": "RPV - O1H", "CA": "REV - O1H", "CF": "SILS - Yes/No",
}
DEFAULTS = {
    "F": "P100", "G": "French", "H": "English", "I": "Family", "J": "Function",
    "M": "Part", "N": "Route", "O": 1, "P": 1, "Q": "Common to all",
    "Y": "Seller", "Z": "Seller company", "AA": "Shipper", "AB": "France",
    "AF": "Pack remark", "AU": "UC1", "AV": "Box", "AY": "Type A",
    "AZ": 10, "BA": 20, "BB": 30, "BC": "Standard", "BD": "Reusable",
    "BE": "Steel", "BO": 0.5, "BT": 0.5, "BV": 80, "BW": 75,
    "BX": 1, "BY": 2, "CA": 3, "CF": "No",
}


def make_workbook(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    for letter, value in HEADERS.items():
        ws[f"{letter}4"] = value
    for index, overrides in enumerate(rows):
        excel_row = 7 + index
        values = {**DEFAULTS, **overrides}
        for letter, value in values.items():
            ws[f"{letter}{excel_row}"] = value
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


@pytest.fixture(autouse=True)
def load_rules():
    T._STATE.update(key=None, bytes=None, wb=None, tables={})
    R.set_rules_file(RULES_PATH.read_bytes())
    yield
    T._STATE.update(key=None, bytes=None, wb=None, tables={})


def check(rows):
    T.set_file(make_workbook(rows))
    return R.check_all_rules(SHEET)


def failures(result, rule_id):
    return sorted(x["cell"] for x in result["findings"] if x["rule_id"] == rule_id)


def test_loads_all_31_rules():
    result = check([{}])
    assert result["rules_loaded"] == 31


def test_duplicate_part_number_flags_every_occurrence_case_and_space_insensitive():
    result = check([{"F": " P100 "}, {"F": "p100"}, {"F": "P200"}])
    assert failures(result, "R001") == ["F7", "F8"]


def test_blank_required_cell_is_flagged():
    result = check([{"G": "   "}])
    assert failures(result, "R002") == ["G7"]


def test_sils_yes_allows_blank_uc_code():
    result = check([{"AU": None, "CF": "Yes"}])
    assert failures(result, "R015") == []
    assert next(x for x in result["rule_summary"] if x["rule_id"] == "R016")["status"] == "PASS"


def test_sils_no_requires_uc_code():
    result = check([{"AU": None, "CF": "No"}])
    assert failures(result, "R015") == ["AU7"]


def test_zero_in_O_requires_unique_to_O1H():
    result = check([{"O": 0, "Q": "Common to all"}])
    assert failures(result, "R008") == ["Q7"]
    clean = check([{"O": 0, "Q": "Unique to O1H"}])
    assert failures(clean, "R008") == []


def test_zero_in_P_requires_unique_to_P1H():
    result = check([{"P": 0, "Q": "Common to all"}])
    assert failures(result, "R009") == ["Q7"]
    clean = check([{"P": 0, "Q": "Unique to P1H"}])
    assert failures(clean, "R009") == []


def test_P1H_and_O1H_required_fields():
    result = check([{"O": 1, "P": 0, "BO": None, "BT": None, "BV": None, "BW": None}])
    assert failures(result, "R025") == ["BO7"]
    assert failures(result, "R027") == ["BV7"]
    assert failures(result, "R026") == []
    assert failures(result, "R028") == []


def test_whole_number_rules_reject_decimal_and_text():
    result = check([{"BX": 1.5, "BY": "N/A", "CA": 2}])
    assert failures(result, "R029") == ["BX7"]
    assert failures(result, "R030") == ["BY7"]
    assert failures(result, "R031") == []


def test_clean_row_has_no_cell_level_failures():
    result = check([{}])
    assert result["findings"] == []


def test_rule_results_include_pass_fail_review_statuses():
    result = check([{"O": 0, "Q": "Common to all", "BX": 1.5}])
    statuses = {x["status"] for x in result["rule_summary"]}
    assert "PASS" in statuses
    assert "FAIL" in statuses
    assert "REVIEW" in statuses or result["rules_needing_review"] == 0
